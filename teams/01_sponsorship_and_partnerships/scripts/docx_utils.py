import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Color Palette Constants
HEX_NAVY = "0A2540"       # Primary Dark Navy
HEX_SLATE_BLUE = "1E3A8A" # Secondary Header Blue
HEX_AMBER = "D97706"      # Accent Amber / Gold
HEX_DARK_TEXT = "1E293B"  # Charcoal Body Text
HEX_MUTED = "64748B"      # Muted Slate
HEX_ZEBRA = "F8FAFC"      # Table Alternating Row
HEX_BORDER = "CBD5E1"     # Table Borders
HEX_BLUE_BG = "EFF6FF"    # Info Callout BG
HEX_AMBER_BG = "FEF3C7"   # Warning/Highlight Callout BG
HEX_GREEN_BG = "ECFDF5"   # Success/Finance Callout BG

COLOR_NAVY = RGBColor(10, 37, 64)
COLOR_SLATE_BLUE = RGBColor(30, 58, 138)
COLOR_AMBER = RGBColor(217, 119, 6)
COLOR_DARK_TEXT = RGBColor(30, 41, 59)
COLOR_MUTED = RGBColor(100, 116, 139)

def init_document(title_short="KAIROS 2027 SPONSORSHIP"):
    doc = docx.Document()
    
    # Page Margins: 0.75 in
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        section.different_first_page_header_footer = False
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"AWS STUDENT BUILDER GROUP — SIST CHAPTER  |  {title_short}")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = COLOR_MUTED
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun1 = fp.add_run("KAIROS 2027 (Grand Edition)  •  Official Sponsorship Operations  •  Confidential & Internal Use")
        frun1.font.name = "Arial"
        frun1.font.size = Pt(8)
        frun1.font.color.rgb = COLOR_MUTED
        
    # Set default style font
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Arial'
    font.size = Pt(10)
    font.color.rgb = COLOR_DARK_TEXT
    
    return doc

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_header_banner(doc, title, subtitle, badges=None):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, HEX_NAVY)
    set_cell_margins(cell, top=200, bottom=200, left=240, right=240)
    
    # Left amber border accent
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="48" w:space="0" w:color="{HEX_AMBER}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    
    # Org Subtitle
    run_org = p.add_run("SATHYABAMA INSTITUTE OF SCIENCE AND TECHNOLOGY  •  AWS STUDENT BUILDER GROUP\n")
    run_org.font.name = "Arial"
    run_org.font.size = Pt(8.5)
    run_org.font.bold = True
    run_org.font.color.rgb = RGBColor(245, 158, 11) # Amber
    
    # Main Title
    run_title = p.add_run(f"{title}\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(17)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(255, 255, 255)
    
    # Subtitle
    run_sub = p.add_run(subtitle)
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(226, 232, 240)
    
    if badges:
        run_badges = p.add_run("\n\n" + "  |  ".join(badges))
        run_badges.font.name = "Arial"
        run_badges.font.size = Pt(8.5)
        run_badges.font.bold = True
        run_badges.font.color.rgb = RGBColor(253, 230, 138)
        
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(8)

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = COLOR_NAVY
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = COLOR_SLATE_BLUE
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_body(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(9.5)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_bullet(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2.5)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix + ": ")
        r_bold.font.name = "Arial"
        r_bold.font.size = Pt(9.5)
        r_bold.font.bold = True
        r_bold.font.color.rgb = COLOR_DARK_TEXT
        
    r_text = p.add_run(text)
    r_text.font.name = "Arial"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = COLOR_DARK_TEXT
    return p

def add_callout(doc, text, title=None, style="info"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    if style == "info":
        bg_color = HEX_BLUE_BG
        border_color = "2563EB"
        title_color = RGBColor(30, 58, 138)
    elif style == "warning":
        bg_color = HEX_AMBER_BG
        border_color = HEX_AMBER
        title_color = RGBColor(180, 83, 9)
    else: # success / money
        bg_color = HEX_GREEN_BG
        border_color = "16A34A"
        title_color = RGBColor(22, 101, 52)
        
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="36" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    if title:
        rt = p.add_run(f"★ {title}\n")
        rt.font.name = "Arial"
        rt.font.size = Pt(10)
        rt.font.bold = True
        rt.font.color.rgb = title_color
        
    rx = p.add_run(text)
    rx.font.name = "Arial"
    rx.font.size = Pt(9)
    rx.font.color.rgb = COLOR_DARK_TEXT
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(5)

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].text = header_text
        set_cell_background(hdr_cells[i], HEX_NAVY)
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for row_idx, row_data in enumerate(data):
        row_cells = table.rows[row_idx + 1].cells
        is_even = (row_idx % 2 == 1)
        row_bg = HEX_ZEBRA if is_even else "FFFFFF"
        
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_background(row_cells[col_idx], row_bg)
            set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Arial"
                run.font.size = Pt(8.5)
                run.font.color.rgb = COLOR_DARK_TEXT
                
    # Set thin borders on all cells
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
            
    # Apply column widths if provided
    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = width
                
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    
    return table
